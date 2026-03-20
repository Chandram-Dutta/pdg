from __future__ import annotations

import os
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from pdg.models import ReportDef, SectionDef, SectionState
from pdg.plotting import plot_section


def generate_report(
    report: ReportDef,
    section_states: dict[str, SectionState],
    output_path: Path,
    template_path: Path | None = None,
) -> None:
    """Build a Word document from a report definition and CSV data."""
    doc = Document(template_path) if template_path else Document()

    for section in report.sections:
        state = section_states.get(section.key)
        if not state or not state.csv_files:
            continue
        _write_section(doc, section, state)

    doc.save(output_path)


def _write_section(
    doc: Document,
    section: SectionDef,
    state: SectionState,
) -> None:
    doc.add_heading(section.title, level=1)

    for idx, csv_path in enumerate(state.csv_files):
        heading = _heading_for(section, idx)

        df = pd.read_csv(csv_path)

        _add_table(doc, df, title=f"Table {idx + 1}: {heading}")

        img = plot_section(df, heading, section)
        if img:
            doc.add_picture(img, width=Inches(6))
            os.remove(img)

    if state.analysis_approved and state.analysis_text:
        doc.add_heading("Analysis", level=2)
        for para in state.analysis_text.split("\n\n"):
            if para.strip():
                doc.add_paragraph(para.strip())


def _heading_for(section: SectionDef, idx: int) -> str:
    if section.loading_levels:
        label = section.loading_levels[idx] if idx < len(section.loading_levels) else f"#{idx + 1}"
        return f"{section.heading_prefix} {label} Loading"
    return section.heading_prefix


def _add_table(doc: Document, df: pd.DataFrame, title: str | None = None) -> None:
    col_count = max(len(df.columns), 1)
    start_rows = 2 if title else 1
    table = doc.add_table(rows=start_rows, cols=col_count)
    table.style = "Table Grid"
    table.allow_autofit = True

    header_row_idx = 0
    if title:
        title_row = table.rows[0]
        title_cell = title_row.cells[0]
        if col_count > 1:
            title_cell = title_cell.merge(title_row.cells[-1])
        title_para = title_cell.paragraphs[0]
        title_run = title_para.add_run(title)
        title_run.bold = True
        title_run.font.size = Pt(11)
        title_para.paragraph_format.keep_with_next = True
        title_para.paragraph_format.space_after = Pt(2)
        header_row_idx = 1

    for ci, col in enumerate(df.columns):
        hdr_cell = table.rows[header_row_idx].cells[ci]
        hdr_cell.text = str(col)
        hdr_para = hdr_cell.paragraphs[0]
        if hdr_para.runs:
            hdr_para.runs[0].bold = True
        hdr_para.paragraph_format.keep_with_next = True
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for ci, val in enumerate(row):
            cells[ci].text = str(val)
    # Allow rows to split across pages so tables flow naturally
    for row in table.rows:
        tr = row._tr
        tr_pr = tr.get_or_add_trPr()
        cant_split = tr_pr.find(qn("w:cantSplit"))
        if cant_split is not None:
            tr_pr.remove(cant_split)
